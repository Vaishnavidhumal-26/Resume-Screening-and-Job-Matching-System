
import re
import os
import tempfile
import streamlit as st
import numpy as np
import pymupdf

from docx import Document
from sentence_transformers import SentenceTransformer
from sklearn.metrics.pairwise import cosine_similarity

# Page settings
st.set_page_config(
    page_title="Resume Job Matcher",
    page_icon="📄",
    layout="wide"
)

st.title("📄 Resume Screening and Job Matching")
st.write(
    "Upload your resume and enter a job description "
    "to analyze skill matching and semantic similarity."
)

# Skills list
skill_keywords = [
    "python", "sql", "mysql", "excel", "power bi", "tableau",
    "pandas", "numpy", "statistics", "data analysis",
    "data cleaning", "data visualization", "machine learning",
    "deep learning", "generative ai", "prompt engineering",
    "llm", "llms", "nlp", "natural language processing",
    "artificial intelligence", "neural networks",
    "model evaluation", "data preprocessing", "rag",
    "transformers", "pytorch", "tensorflow", "scikit-learn",
    "java", "html", "css", "javascript", "communication",
    "customer service", "sales", "github", "powerpoint",
    "problem solving", "vlookup", "xlookup"
]

@st.cache_resource
def load_model():
    return SentenceTransformer(
        "sentence-transformers/all-MiniLM-L6-v2"
    )

def clean_text(text):
    text = text.lower()
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"[^a-z0-9+#.\s]", "", text)
    return text.strip()

def extract_skills(text):
    text = text.lower()
    found = []

    for skill in skill_keywords:
        pattern = r"(?<!\w)" + re.escape(skill) + r"(?!\w)"
        if re.search(pattern, text):
            found.append(skill)

    return sorted(set(found))

def extract_pdf_text(file_path):
    text = ""

    with pymupdf.open(file_path) as pdf:
        for page in pdf:
            text += page.get_text() + "\n"

    return text.strip()

def extract_docx_text(file_path):
    document = Document(file_path)
    parts = [
        p.text for p in document.paragraphs if p.text.strip()
    ]

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)

    return "\n".join(parts).strip()

def extract_resume_text(file_path):
    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".pdf":
        return extract_pdf_text(file_path)
    elif extension == ".docx":
        return extract_docx_text(file_path)

    raise ValueError("Please upload a PDF or DOCX resume.")

# Upload resume
uploaded_file = st.file_uploader(
    "Upload Resume (PDF or DOCX)",
    type=["pdf", "docx"]
)

# Job description
job_description = st.text_area(
    "Paste Job Description",
    height=200,
    placeholder="Paste the job title, required skills and responsibilities here..."
)

if st.button("Analyze Resume", type="primary"):
    if uploaded_file is None:
        st.error("Please upload your resume first.")
    elif not job_description.strip():
        st.error("Please enter a job description.")
    else:
        try:
            with tempfile.NamedTemporaryFile(
                delete=False,
                suffix=os.path.splitext(uploaded_file.name)[1]
            ) as temp_file:
                temp_file.write(uploaded_file.getvalue())
                temp_path = temp_file.name

            try:
                resume_text = extract_resume_text(temp_path)
            finally:
                os.remove(temp_path)

            if not resume_text.strip():
                st.error(
                    "No text found. If your PDF is scanned, "
                    "OCR may be required."
                )
                st.stop()

            clean_resume = clean_text(resume_text)
            clean_job = clean_text(job_description)

            resume_skills = set(extract_skills(clean_resume))
            required_skills = set(extract_skills(clean_job))

            matched_skills = sorted(
                resume_skills & required_skills
            )
            missing_skills = sorted(
                required_skills - resume_skills
            )

            if required_skills:
                skill_score = round(
                    len(matched_skills) /
                    len(required_skills) * 100, 2
                )
            else:
                skill_score = 0.0

            with st.spinner("Loading model and analyzing text..."):
                model = load_model()

                embeddings = model.encode(
                    [clean_resume, clean_job],
                    convert_to_numpy=True
                )

                semantic_score = round(
                    float(cosine_similarity(
                        [embeddings[0]], [embeddings[1]]
                    )[0][0]) * 100,
                    2
                )

            combined_score = round(
                0.5 * skill_score + 0.5 * semantic_score,
                2
            )

            st.success("Resume analysis completed!")

            st.subheader("Matching Results")

            col1, col2, col3 = st.columns(3)
            col1.metric("Skill Match", f"{skill_score}%")
            col2.metric(
                "Semantic Similarity", f"{semantic_score}%"
            )
            col3.metric(
                "Combined Match Score", f"{combined_score}%"
            )

            st.progress(
                max(0.0, min(combined_score / 100, 1.0))
            )

            col1, col2 = st.columns(2)

            with col1:
                st.subheader("✅ Matched Skills")
                if matched_skills:
                    for skill in matched_skills:
                        st.write(f"- {skill}")
                else:
                    st.write("No matching skills found.")

            with col2:
                st.subheader("📌 Skills to Develop")
                if missing_skills:
                    for skill in missing_skills:
                        st.write(f"- {skill}")
                else:
                    st.write("No missing skills found in the keyword list.")

            with st.expander("View extracted resume text"):
                st.text(resume_text)

            st.caption(
                "Note: These scores are experimental matching "
                "indicators, not hiring probabilities or validated "
                "model accuracy. Results depend on the skill keyword "
                "list and the resume text extracted."
            )

        except Exception as error:
            st.error(f"Something went wrong: {error}")
